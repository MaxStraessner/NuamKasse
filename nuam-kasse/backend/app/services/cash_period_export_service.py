from collections import defaultdict
from datetime import datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload

from app.models.cashbook import Cashbook
from app.models.cash_period import CashPeriod, CashPeriodStatus
from app.models.category import Category
from app.models.category import CategoryType
from app.models.expense import Expense
from app.services.cash_summary_service import get_cash_period_summary


class CashPeriodExportError(ValueError):
    pass


TITLE_FILL = PatternFill("solid", fgColor="173F5F")
HEADER_FILL = PatternFill("solid", fgColor="20639B")
SUBTLE_FILL = PatternFill("solid", fgColor="EAF2F8")
TOTAL_FILL = PatternFill("solid", fgColor="D6EAF8")
WHITE_FONT = Font(color="FFFFFF", bold=True)
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN_GRAY = Side(style="thin", color="D5D8DC")


def _money_format(currency: str) -> str:
    return f'#,##0.00 "{currency}";[Red]-#,##0.00 "{currency}"'


def _as_number(value: str | Decimal) -> float:
    return float(Decimal(str(value)))


def _plain_datetime(value: datetime | None) -> datetime | None:
    if value is None:
        return None
    return value.replace(tzinfo=None) if value.tzinfo else value


def _period_range(cash_period: CashPeriod) -> str:
    end_date = cash_period.end_date.strftime("%d.%m.%Y") if cash_period.end_date else "heute"
    return f"{cash_period.start_date:%d.%m.%Y} – {end_date}"


def _style_title(sheet, title: str, end_column: int) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
    cell = sheet.cell(1, 1, title)
    cell.fill = TITLE_FILL
    cell.font = Font(color="FFFFFF", bold=True, size=18)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 30


def _style_header(row) -> None:
    for cell in row:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = Border(bottom=THIN_GRAY)


def _fit_columns(sheet, widths: dict[int, float]) -> None:
    for column_index, width in widths.items():
        sheet.column_dimensions[get_column_letter(column_index)].width = width


def _add_table(sheet, *, name: str, start_row: int, end_row: int, end_column: int) -> None:
    if end_row <= start_row:
        return
    table = Table(
        displayName=name,
        ref=f"A{start_row}:{get_column_letter(end_column)}{end_row}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def _load_expenses(db: Session, cash_period_id: int) -> list[Expense]:
    return list(
        db.scalars(
            select(Expense)
            .options(
                joinedload(Expense.category).joinedload(Category.parent),
                joinedload(Expense.created_by),
                joinedload(Expense.voided_by),
            )
            .where(Expense.cash_period_id == cash_period_id)
            .order_by(Expense.created_at.asc(), Expense.id.asc())
        )
    )


def _category_analytics(expenses: list[Expense]) -> list[dict[str, object]]:
    aggregates: dict[tuple[str, str], dict[str, object]] = defaultdict(
        lambda: {
            "income": Decimal("0.00"),
            "expense": Decimal("0.00"),
            "count": 0,
        }
    )
    for expense in expenses:
        if expense.is_voided:
            continue
        parent = expense.category.parent
        root_name = parent.name if parent else expense.category.name
        child_name = expense.category.name if parent else "—"
        item = aggregates[(root_name, child_name)]
        if expense.transaction_type == CategoryType.income:
            item["income"] += expense.amount
        else:
            item["expense"] += expense.amount
        item["count"] += 1
    return [
        {
            "root": root,
            "child": child,
            "income": values["income"],
            "expense": values["expense"],
            "net": values["income"] - values["expense"],
            "count": values["count"],
        }
        for (root, child), values in sorted(aggregates.items())
    ]


def _user_analytics(expenses: list[Expense]) -> list[dict[str, object]]:
    aggregates: dict[str, dict[str, object]] = defaultdict(
        lambda: {
            "income": Decimal("0.00"),
            "expense": Decimal("0.00"),
            "count": 0,
        }
    )
    for expense in expenses:
        if expense.is_voided:
            continue
        item = aggregates[expense.created_by.display_name]
        if expense.transaction_type == CategoryType.income:
            item["income"] += expense.amount
        else:
            item["expense"] += expense.amount
        item["count"] += 1
    return [
        {
            "name": name,
            "income": values["income"],
            "expense": values["expense"],
            "net": values["income"] - values["expense"],
            "count": values["count"],
        }
        for name, values in sorted(aggregates.items())
    ]


def build_cash_period_export(
    db: Session,
    *,
    cashbook: Cashbook,
    cash_period: CashPeriod,
) -> BytesIO:
    summary = get_cash_period_summary(db, cash_period)
    expenses = _load_expenses(db, cash_period.id)
    categories = _category_analytics(expenses)
    users = _user_analytics(expenses)
    currency = cash_period.currency
    money_format = _money_format(currency)

    workbook = Workbook()
    overview = workbook.active
    overview.title = "Übersicht"
    _style_title(overview, "Nuam Kasse", 12)
    status_label = (
        "Laufende Kasse"
        if cash_period.status == CashPeriodStatus.active
        else "Abgeschlossen"
    )
    metadata = (
        ("Kasse", cashbook.name),
        ("Bezeichnung", cash_period.name),
        ("Zeitraum", _period_range(cash_period)),
        ("Status", status_label),
        ("Abgeschlossen am", _plain_datetime(cash_period.closed_at)),
        ("Anzahl Buchungen", summary["active_expense_count"]),
    )
    for row_index, (label, value) in enumerate(metadata, 3):
        overview.cell(row_index, 1, label).font = Font(bold=True, color="173F5F")
        overview.cell(row_index, 2, value)
    overview["B7"].number_format = "dd.mm.yyyy hh:mm"

    overview["A9"] = "Kennzahl"
    overview["B9"] = "Wert"
    _style_header(overview[9][0:2])
    metrics = (
        ("Anfangsbestand", summary["opening_amount"]),
        ("Gesamteinnahmen", summary["income_amount"]),
        ("Gesamtausgaben", summary["spent_amount"]),
        ("Saldo", summary["net_amount"]),
        ("Endbestand", summary["remaining_amount"]),
        ("Buchungen", summary["active_expense_count"]),
    )
    for row_index, (label, value) in enumerate(metrics, 10):
        overview.cell(row_index, 1, label)
        overview.cell(row_index, 2, _as_number(value) if label != "Buchungen" else value)
        if label != "Buchungen":
            overview.cell(row_index, 2).number_format = money_format
    overview["A15"].fill = TOTAL_FILL
    overview["B15"].fill = TOTAL_FILL
    overview["A15"].font = Font(bold=True)
    overview["B15"].font = Font(bold=True)

    overview["G3"] = "Typ"
    overview["H3"] = "Betrag"
    _style_header(overview[3][6:8])
    overview["G4"] = "Einnahmen"
    overview["H4"] = _as_number(summary["income_amount"])
    overview["G5"] = "Ausgaben"
    overview["H5"] = _as_number(summary["spent_amount"])
    overview["H4"].number_format = money_format
    overview["H5"].number_format = money_format
    if Decimal(str(summary["income_amount"])) or Decimal(str(summary["spent_amount"])):
        comparison_chart = BarChart()
        comparison_chart.title = "Einnahmen und Ausgaben im Vergleich"
        comparison_chart.y_axis.title = currency
        comparison_chart.height = 7
        comparison_chart.width = 12
        comparison_chart.add_data(Reference(overview, min_col=8, min_row=3, max_row=5), titles_from_data=True)
        comparison_chart.set_categories(Reference(overview, min_col=7, min_row=4, max_row=5))
        overview.add_chart(comparison_chart, "D8")

    expense_by_root: dict[str, Decimal] = defaultdict(lambda: Decimal("0.00"))
    for item in categories:
        amount = Decimal(str(item["expense"]))
        if amount > 0:
            expense_by_root[str(item["root"])] += amount
    chart_start = 18
    overview.cell(chart_start, 1, "Oberkategorie")
    overview.cell(chart_start, 2, "Ausgaben")
    _style_header(overview[chart_start][0:2])
    for row_index, (name, amount) in enumerate(
        sorted(expense_by_root.items(), key=lambda item: item[1], reverse=True),
        chart_start + 1,
    ):
        overview.cell(row_index, 1, name)
        overview.cell(row_index, 2, _as_number(amount)).number_format = money_format
    if expense_by_root:
        expense_chart = PieChart()
        expense_chart.title = "Ausgaben nach Oberkategorie"
        expense_chart.height = 8
        expense_chart.width = 12
        expense_chart.add_data(
            Reference(overview, min_col=2, min_row=chart_start, max_row=chart_start + len(expense_by_root)),
            titles_from_data=True,
        )
        expense_chart.set_categories(
            Reference(overview, min_col=1, min_row=chart_start + 1, max_row=chart_start + len(expense_by_root))
        )
        overview.add_chart(expense_chart, "D24")

    daily: dict[object, dict[str, Decimal]] = defaultdict(
        lambda: {"income": Decimal("0.00"), "expense": Decimal("0.00")}
    )
    for expense in expenses:
        if expense.is_voided:
            continue
        key = expense.created_at.date()
        target = "income" if expense.transaction_type == CategoryType.income else "expense"
        daily[key][target] += expense.amount
    if len(daily) >= 2:
        overview["J3"] = "Datum"
        overview["K3"] = "Einnahmen"
        overview["L3"] = "Ausgaben"
        _style_header(overview[3][9:12])
        for row_index, (day, amounts) in enumerate(sorted(daily.items()), 4):
            overview.cell(row_index, 10, day).number_format = "dd.mm.yyyy"
            overview.cell(row_index, 11, _as_number(amounts["income"])).number_format = money_format
            overview.cell(row_index, 12, _as_number(amounts["expense"])).number_format = money_format
        trend_chart = LineChart()
        trend_chart.title = "Entwicklung der Kasse"
        trend_chart.y_axis.title = currency
        trend_chart.height = 8
        trend_chart.width = 12
        trend_chart.add_data(
            Reference(overview, min_col=11, max_col=12, min_row=3, max_row=3 + len(daily)),
            titles_from_data=True,
        )
        trend_chart.set_categories(
            Reference(overview, min_col=10, min_row=4, max_row=3 + len(daily))
        )
        overview.add_chart(trend_chart, "D40")

    user_start = chart_start + max(len(expense_by_root), 1) + 3
    overview.cell(user_start, 1, "Erfasst von")
    overview.cell(user_start, 2, "Einnahmen")
    overview.cell(user_start, 3, "Ausgaben")
    overview.cell(user_start, 4, "Saldo")
    overview.cell(user_start, 5, "Buchungen")
    _style_header(overview[user_start][0:5])
    for row_index, item in enumerate(users, user_start + 1):
        overview.cell(row_index, 1, item["name"])
        overview.cell(row_index, 2, _as_number(item["income"])).number_format = money_format
        overview.cell(row_index, 3, _as_number(item["expense"])).number_format = money_format
        overview.cell(row_index, 4, _as_number(item["net"])).number_format = money_format
        overview.cell(row_index, 5, item["count"])
    overview.freeze_panes = "A9"
    overview.sheet_view.showGridLines = False
    overview.print_title_rows = "1:9"
    overview.page_setup.orientation = "landscape"
    overview.page_setup.fitToWidth = 1
    overview.sheet_properties.pageSetUpPr.fitToPage = True
    _fit_columns(overview, {1: 24, 2: 22, 3: 16, 4: 16, 5: 14, 7: 18, 8: 18, 10: 14, 11: 18, 12: 18})

    category_sheet = workbook.create_sheet("Kategorien")
    _style_title(category_sheet, "Kategorienanalyse", 6)
    category_headers = ("Oberkategorie", "Unterkategorie", "Einnahmen", "Ausgaben", "Saldo", "Buchungen")
    for column_index, header in enumerate(category_headers, 1):
        category_sheet.cell(3, column_index, header)
    _style_header(category_sheet[3][0:6])
    for row_index, item in enumerate(categories, 4):
        category_sheet.cell(row_index, 1, item["root"])
        category_sheet.cell(row_index, 2, item["child"])
        category_sheet.cell(row_index, 3, _as_number(item["income"])).number_format = money_format
        category_sheet.cell(row_index, 4, _as_number(item["expense"])).number_format = money_format
        category_sheet.cell(row_index, 5, _as_number(item["net"])).number_format = money_format
        category_sheet.cell(row_index, 6, item["count"])
    category_total_row = max(4, 4 + len(categories))
    category_sheet.cell(category_total_row, 1, "Gesamt")
    category_sheet.cell(category_total_row, 3, _as_number(summary["income_amount"]))
    category_sheet.cell(category_total_row, 4, _as_number(summary["spent_amount"]))
    category_sheet.cell(category_total_row, 5, _as_number(summary["net_amount"]))
    category_sheet.cell(category_total_row, 6, summary["active_expense_count"])
    for cell in category_sheet[category_total_row][0:6]:
        cell.fill = TOTAL_FILL
        cell.font = Font(bold=True)
    for column in range(3, 6):
        category_sheet.cell(category_total_row, column).number_format = money_format
    _add_table(
        category_sheet,
        name="KategorienAnalyse",
        start_row=3,
        end_row=3 + len(categories),
        end_column=6,
    )
    category_sheet.freeze_panes = "A4"
    category_sheet.auto_filter.ref = f"A3:F{max(3, 3 + len(categories))}"
    category_sheet.sheet_view.showGridLines = False
    category_sheet.page_setup.orientation = "landscape"
    category_sheet.page_setup.fitToWidth = 1
    category_sheet.sheet_properties.pageSetUpPr.fitToPage = True
    _fit_columns(category_sheet, {1: 26, 2: 28, 3: 18, 4: 18, 5: 18, 6: 14})

    transaction_sheet = workbook.create_sheet("Buchungen")
    _style_title(transaction_sheet, "Buchungen", 11)
    transaction_headers = (
        "Datum",
        "Uhrzeit",
        "Art",
        "Kategorie",
        "Unterkategorie",
        "Beschreibung",
        "Betrag",
        "Erfasst von",
        "Status",
        "Storniert am",
        "Stornierungsgrund",
    )
    for column_index, header in enumerate(transaction_headers, 1):
        transaction_sheet.cell(3, column_index, header)
    _style_header(transaction_sheet[3][0:11])
    for row_index, expense in enumerate(expenses, 4):
        parent = expense.category.parent
        created_at = _plain_datetime(expense.created_at)
        transaction_sheet.cell(row_index, 1, created_at.date()).number_format = "dd.mm.yyyy"
        transaction_sheet.cell(row_index, 2, created_at.time()).number_format = "hh:mm"
        transaction_sheet.cell(row_index, 3, "Einnahme" if expense.transaction_type == CategoryType.income else "Ausgabe")
        transaction_sheet.cell(row_index, 4, parent.name if parent else expense.category.name)
        transaction_sheet.cell(row_index, 5, expense.category.name if parent else "—")
        transaction_sheet.cell(row_index, 6, expense.note or "")
        transaction_sheet.cell(row_index, 7, _as_number(expense.amount)).number_format = money_format
        transaction_sheet.cell(row_index, 8, expense.created_by.display_name)
        transaction_sheet.cell(row_index, 9, "Storniert" if expense.is_voided else "Gültig")
        transaction_sheet.cell(row_index, 10, _plain_datetime(expense.voided_at)).number_format = "dd.mm.yyyy hh:mm"
        transaction_sheet.cell(row_index, 11, expense.void_reason or "")
    _add_table(
        transaction_sheet,
        name="BuchungenListe",
        start_row=3,
        end_row=3 + len(expenses),
        end_column=11,
    )
    transaction_sheet.freeze_panes = "A4"
    transaction_sheet.auto_filter.ref = f"A3:K{max(3, 3 + len(expenses))}"
    transaction_sheet.sheet_view.showGridLines = False
    transaction_sheet.page_setup.orientation = "landscape"
    transaction_sheet.page_setup.fitToWidth = 1
    transaction_sheet.sheet_properties.pageSetUpPr.fitToPage = True
    _fit_columns(
        transaction_sheet,
        {1: 14, 2: 10, 3: 14, 4: 24, 5: 26, 6: 34, 7: 18, 8: 22, 9: 14, 10: 20, 11: 30},
    )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
