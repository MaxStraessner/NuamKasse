from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from app.models import CashPeriod, CashPeriodStatus, Category, CategoryType, Expense, UserRole
from conftest import create_test_user, get_test_cashbook


def login(client, username: str):
    return client.post("/api/v1/auth/login", json={"username": username, "password": "password-123"})


def test_closed_period_export_contains_analytics_charts_categories_and_all_bookings(client, db_session):
    admin = create_test_user(db_session, username="admin", role=UserRole.admin)
    member = create_test_user(db_session, username="nuam", role=UserRole.member)
    cashbook = get_test_cashbook(db_session)
    period = CashPeriod(
        cashbook_id=cashbook.id,
        name="Juli 2026",
        opening_amount=Decimal("1000.00"),
        currency="THB",
        start_date=date(2026, 7, 1),
        end_date=date(2026, 7, 31),
        status=CashPeriodStatus.closed,
        created_by_user_id=admin.id,
        closed_by_user_id=admin.id,
        closed_at=datetime(2026, 7, 31, 18, 0, tzinfo=timezone.utc),
    )
    expense_root = Category(
        cashbook_id=cashbook.id,
        user_id=admin.id,
        name="Lebensmittel",
        name_normalized="lebensmittel",
        icon_key="utensils",
        color_key="orange",
        category_type=CategoryType.expense,
        sort_order=1,
        is_active=True,
    )
    income_root = Category(
        cashbook_id=cashbook.id,
        user_id=admin.id,
        name="Einnahmen",
        name_normalized="einnahmen",
        icon_key="wallet",
        color_key="green",
        category_type=CategoryType.income,
        sort_order=2,
        is_active=True,
    )
    db_session.add_all([period, expense_root, income_root])
    db_session.flush()
    child = Category(
        cashbook_id=cashbook.id,
        user_id=admin.id,
        name="Markt",
        name_normalized="markt",
        icon_key="shopping-cart",
        color_key="orange",
        category_type=CategoryType.expense,
        parent_category_id=expense_root.id,
        sort_order=1,
        is_active=True,
    )
    db_session.add(child)
    db_session.flush()
    db_session.add_all(
        [
            Expense(
                cash_period_id=period.id,
                category_id=child.id,
                amount=Decimal("250.00"),
                transaction_type=CategoryType.expense,
                currency="THB",
                created_by_user_id=member.id,
                note="Wochenmarkt",
                created_at=datetime(2026, 7, 5, 9, 30, tzinfo=timezone.utc),
            ),
            Expense(
                cash_period_id=period.id,
                category_id=income_root.id,
                amount=Decimal("500.00"),
                transaction_type=CategoryType.income,
                currency="THB",
                created_by_user_id=admin.id,
                note="Einlage",
                created_at=datetime(2026, 7, 6, 10, 0, tzinfo=timezone.utc),
            ),
        ]
    )
    db_session.commit()

    login(client, "admin")
    response = client.get(f"/api/v1/cash-periods/{period.id}/export.xlsx")
    repeated_response = client.get(f"/api/v1/cash-periods/{period.id}/export.xlsx")

    assert response.status_code == 200
    assert repeated_response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(BytesIO(response.content), data_only=False)
    repeated_workbook = load_workbook(BytesIO(repeated_response.content), data_only=False)
    assert workbook.sheetnames == ["Übersicht", "Kategorien", "Buchungen"]
    assert repeated_workbook.sheetnames == workbook.sheetnames
    assert len(workbook["Übersicht"]._charts) >= 2
    chart_references = " ".join(
        series.val.numRef.f
        for chart in workbook["Übersicht"]._charts
        for series in chart.ser
        if series.val is not None and series.val.numRef is not None
    )
    assert "'Übersicht'!$H$4:$H$5" in chart_references
    assert "'Übersicht'!$B$19" in chart_references
    assert "'Übersicht'!$K$4:$K$5" in chart_references
    assert "'Übersicht'!$L$4:$L$5" in chart_references
    assert workbook["Übersicht"]["A1"].value == "Nuam Kasse"
    metric_values = {
        workbook["Übersicht"].cell(row, 1).value: workbook["Übersicht"].cell(row, 2).value
        for row in range(10, 16)
    }
    assert metric_values["Gesamteinnahmen"] == 500
    assert metric_values["Gesamtausgaben"] == 250
    assert metric_values["Saldo"] == 250
    assert metric_values["Endbestand"] == 1250

    category_rows = list(workbook["Kategorien"].iter_rows(min_row=4, values_only=True))
    assert ("Lebensmittel", "Markt", 0, 250, -250, 1) in category_rows
    assert ("Einnahmen", "—", 500, 0, 500, 1) in category_rows

    booking_rows = list(workbook["Buchungen"].iter_rows(min_row=4, values_only=True))
    assert len(booking_rows) == 2
    assert {row[7] for row in booking_rows} == {admin.display_name, member.display_name}
    assert {row[5] for row in booking_rows} == {"Wochenmarkt", "Einlage"}
    assert {row[6] for row in booking_rows} == {250, 500}


def test_active_period_can_be_exported(client, db_session):
    admin = create_test_user(db_session, username="admin", role=UserRole.admin)
    period = CashPeriod(
        cashbook_id=get_test_cashbook(db_session).id,
        name="August 2026",
        opening_amount=Decimal("100.00"),
        currency="THB",
        start_date=date(2026, 8, 1),
        status=CashPeriodStatus.active,
        created_by_user_id=admin.id,
    )
    db_session.add(period)
    db_session.commit()
    login(client, "admin")

    response = client.get(f"/api/v1/cash-periods/{period.id}/export.xlsx")

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), data_only=False)
    assert workbook["Übersicht"]["B6"].value == "Laufende Kassenperiode"
    assert workbook.sheetnames == ["Übersicht", "Kategorien", "Buchungen"]
